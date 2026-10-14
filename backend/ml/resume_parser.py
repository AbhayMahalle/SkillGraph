"""
Resume Parser & Text Extraction Service for SkillGraph.
Supports secure parsing of PDF, DOCX, and TXT files with format validation,
magic byte inspection, size limits, and user-isolated file storage.
"""

import io
import os
import re
import uuid
from pathlib import Path
from typing import Dict, Any, Tuple
import pymupdf
import docx

from backend.database.db import REPO_ROOT

UPLOAD_BASE_DIR = REPO_ROOT / "data" / "uploads" / "resumes"
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB limit
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

class ResumeParsingError(Exception):
    """Custom exception raised when resume file parsing fails."""
    pass

def sanitize_filename(filename: str) -> str:
    """Removes path traversal characters and unsafe symbols from filename."""
    base = os.path.basename(filename)
    clean = re.sub(r"[^\w\.\-]", "_", base)
    return clean[:100]

def validate_and_extract_text(filename: str, content: bytes) -> Tuple[str, str]:
    """
    Validates file extension, size, and magic bytes, then extracts plain text.
    Returns: (extracted_text, detected_file_type)
    Raises: ResumeParsingError if invalid, malformed, or unreadable.
    """
    if not content:
        raise ResumeParsingError("Uploaded file is empty (0 bytes).")

    if len(content) > MAX_FILE_SIZE_BYTES:
        raise ResumeParsingError(f"File size ({len(content) / (1024*1024):.1f} MB) exceeds maximum 10 MB limit.")

    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeParsingError(f"Unsupported file extension '{ext}'. Supported formats: PDF (.pdf), Word (.docx), and Text (.txt).")

    extracted_text = ""
    file_type = ""

    # 1. PDF Processing
    if ext == ".pdf":
        if not content.startswith(b"%PDF"):
            raise ResumeParsingError("Invalid PDF file: Missing standard %PDF file header signature.")
        file_type = "PDF"
        try:
            doc = pymupdf.open(stream=content, filetype="pdf")
            if len(doc) == 0:
                doc.close()
                raise ResumeParsingError("PDF document contains 0 pages.")
            
            pages_text = []
            for page in doc:
                pages_text.append(page.get_text())
            doc.close()
            extracted_text = "\n".join(pages_text).strip()
        except Exception as e:
            if isinstance(e, ResumeParsingError):
                raise
            raise ResumeParsingError(f"Unable to read PDF file structure: {str(e)}")

        if not extracted_text:
            raise ResumeParsingError(
                "PDF contains no extractable text. The file appears to be a scanned image or photo without an OCR text layer."
            )

    # 2. DOCX Processing
    elif ext == ".docx":
        if not content.startswith(b"PK\x03\x04"):
            raise ResumeParsingError("Invalid DOCX file: Missing standard ZIP/DOCX package signature.")
        file_type = "DOCX"
        try:
            doc = docx.Document(io.BytesIO(content))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            lines.append(cell.text.strip())
            extracted_text = "\n".join(lines).strip()
        except Exception as e:
            raise ResumeParsingError(f"Unable to read Word (.docx) document structure: {str(e)}")

        if not extracted_text:
            raise ResumeParsingError("Word document (.docx) contains no readable text content.")

    # 3. Plain Text Processing
    elif ext == ".txt":
        file_type = "TXT"
        try:
            extracted_text = content.decode("utf-8").strip()
        except UnicodeDecodeError:
            try:
                extracted_text = content.decode("latin-1").strip()
            except Exception as e:
                raise ResumeParsingError(f"Unable to decode text file: {str(e)}")

        if not extracted_text:
            raise ResumeParsingError("Text file is empty.")

    return extracted_text, file_type

def save_uploaded_resume(user_id: int, filename: str, content: bytes) -> str:
    """
    Saves the validated resume file into a secure, user-isolated directory.
    Returns relative stored file path.
    """
    safe_name = sanitize_filename(filename)
    unique_name = f"{uuid.uuid4().hex[:12]}_{safe_name}"
    user_dir = UPLOAD_BASE_DIR / str(user_id)
    user_dir.mkdir(parents=True, exist_ok=True)

    dest_path = user_dir / unique_name
    with open(dest_path, "wb") as f:
        f.write(content)

    return str(dest_path.relative_to(REPO_ROOT))
